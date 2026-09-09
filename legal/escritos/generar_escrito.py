# -*- coding: utf-8 -*-
"""Escrito de solicitud de permiso de exploración (cateo) — Jáchal."""
import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Arial"
ACC = RGBColor(0x9A, 0x52, 0x30)     # cobre: campos a completar
RED = RGBColor(0xA3, 0x2B, 0x1C)     # plazos fatales
GRY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# ---------- pagina y estilo base ----------
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)      # A4
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(s, m, Cm(2))
USABLE = Cm(17.0)

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(7)
st.paragraph_format.line_spacing = 1.25


def run(p, text, bold=False, italic=False, underline=False, color=None, size=None, caps=False):
    r = p.add_run(text)
    r.bold, r.italic, r.underline = bold, italic, underline
    r.font.name = FONT
    if color is not None:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    if caps:
        r.font.all_caps = True
    return r


def para(runs, align="just", after=7, before=0, line=1.25, first_indent=None, new_page=False):
    p = doc.add_paragraph()
    p.alignment = {"just": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = line
    # Salto antes del párrafo: evita la página en blanco que deja un párrafo
    # vacío con page break cuando el bloque anterior termina cerca del borde.
    p.paragraph_format.page_break_before = new_page
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(first_indent)
    if isinstance(runs, str):
        runs = [(runs, {})]
    for t, f in runs:
        run(p, t, bold=f.get("b"), italic=f.get("i"), underline=f.get("u"),
            color=f.get("c"), size=f.get("s"), caps=f.get("caps"))
    return p


def fill(p, label):
    """Campo a completar: subrayado y en color."""
    return run(p, label, underline=True, color=ACC)


def heading(text):
    return para([(text, {"b": True})], align="left", before=11, after=5)


def bullets(items, size=None):
    for t in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run(p, t, size=size)


def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear"); el.set(qn("w:color"), "auto"); el.set(qn("w:fill"), hexfill)
    tcPr.append(el)


def table(widths_cm, rows, size=9.5, head=True, gap=4):
    t = doc.add_table(rows=0, cols=len(widths_cm))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for ri, cells in enumerate(rows):
        row = t.add_row()
        for ci, spec in enumerate(cells):
            if isinstance(spec, str):
                spec = {"t": spec}
            c = row.cells[ci]
            c.width = Cm(widths_cm[ci])
            p = c.paragraphs[0]
            p.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT,
                           "c": WD_ALIGN_PARAGRAPH.CENTER}[spec.get("a", "l")]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            run(p, spec["t"], bold=spec.get("b") or (head and ri == 0),
                color=spec.get("c"), size=size)
            if head and ri == 0:
                shade(c, "F2F2F2")
    # fijar anchos tambien a nivel columna
    for ci, w in enumerate(widths_cm):
        for cell in t.columns[ci].cells:
            cell.width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(gap)
    return t


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# numero de pagina en el pie
def footer_pagenum():
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.font.size, r.font.name, r.font.color.rgb = Pt(8), FONT, GRY
    for instr in ("begin", "instr", "end"):
        el = OxmlElement("w:fldChar" if instr != "instr" else "w:instrText")
        if instr == "instr":
            el.set(qn("xml:space"), "preserve"); el.text = " PAGE "
        else:
            el.set(qn("w:fldCharType"), instr)
        r._r.append(el)


footer_pagenum()

# ======================= ADVERTENCIA =======================
p = para([("BORRADOR PARA REVISIÓN PROFESIONAL — NO PRESENTAR SIN COMPLETAR",
           {"b": True, "c": ACC, "s": 9.5})], align="left", after=3)
para([("El uso del formulario oficial que provee la Autoridad Minera es obligatorio (art. 8 Ley 688-M). "
       "Este escrito se presenta acompañando ese formulario y volcando en él los mismos datos. Los tramos ", {"s": 9}),
      ("subrayados en color", {"s": 9, "c": ACC, "u": True}),
      (" deben completarse antes de la presentación. La firma corresponde a profesional matriculado.", {"s": 9})],
     after=18, line=1.1)

# ======================= SUMA =======================
para([("SOLICITA PERMISO DE EXPLORACIÓN (CATEO)", {"b": True})], align="right", after=14)
para([("SEÑOR DIRECTOR DE MINERÍA DE LA PROVINCIA DE SAN JUAN:", {"b": True})], align="left", after=12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.first_line_indent = Cm(2)
p.paragraph_format.line_spacing = 1.4
p.paragraph_format.space_after = Pt(9)
fill(p, "[APELLIDO Y NOMBRE / RAZÓN SOCIAL]")
run(p, ", ")
fill(p, "[nacionalidad]")
run(p, ", ")
fill(p, "[estado civil]")
run(p, ", de ")
fill(p, "[edad]")
run(p, " años, de profesión ")
fill(p, "[profesión]")
run(p, ", D.N.I. N.º ")
fill(p, "[documento]")
run(p, ", C.U.I.T. N.º ")
fill(p, "[CUIT]")
run(p, ", con domicilio real en ")
fill(p, "[calle, número, ciudad, provincia]")
run(p, ", constituyendo domicilio legal en ")
fill(p, "[calle y número, Ciudad de San Juan]")
run(p, ", dentro del radio fijado por esa Autoridad Minera, ante V.S. respetuosamente me presento y digo:")

# ======================= I. OBJETO =======================
heading("I. OBJETO")
para([("Que vengo por el presente a solicitar, en los términos de los artículos 25, 27, 28, 29 y 30 del "
       "Código de Minería de la Nación y de los artículos 8, 35, 36 y concordantes de la Ley Provincial "
       "N.º 688-M, el otorgamiento de un ", {}),
      ("PERMISO DE EXPLORACIÓN (CATEO)", {"b": True}),
      (" sobre el área libre que se individualiza en el Capítulo III y cuyas coordenadas se detallan en el ", {}),
      ("Anexo I", {"b": True}), (" del presente.", {})])

# ======================= II. PERSONERIA =======================
heading("II. PERSONERÍA")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.line_spacing = 1.25; p.paragraph_format.space_after = Pt(7)
run(p, "Acredito mi personería mediante ")
fill(p, "[documento de identidad / testimonio de poder / contrato social e inscripción registral]")
run(p, ", cuya copia se acompaña. En caso de actuar por apoderado, se da cumplimiento a lo dispuesto "
       "por el artículo 7 de la Ley N.º 688-M.")

# ======================= III. UBICACION =======================
heading("III. UBICACIÓN Y DETERMINACIÓN DEL ÁREA SOLICITADA")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.line_spacing = 1.25; p.paragraph_format.space_after = Pt(7)
run(p, "El área solicitada se ubica en el departamento ")
run(p, "JÁCHAL", bold=True)
run(p, ", Provincia de San Juan, distrito ")
fill(p, "[distrito]")
run(p, ", paraje ")
fill(p, "[paraje o lugar]")
run(p, ". El punto central del área se sitúa en las coordenadas geográficas latitud 30° 06′ 04″ Sur y "
       "longitud 68° 39′ 18″ Oeste (−30,10116 / −68,65507, WGS84).")

para([("La figura solicitada es un ", {}),
      ("rectángulo de 5,25 km en sentido Este–Oeste por 11,50 km en sentido Norte–Sur", {"b": True}),
      (", cuyos lados observan estrictamente la orientación Norte–Sur y Este–Oeste que imponen el "
       "artículo 25, último párrafo, del Código de Minería y el artículo 38 de la Ley N.º 688-M. "
       "La forma adoptada es la más regular posible, de modo tal que pueda constituirse una pertenencia minera.", {})])

para([("Las coordenadas de los cuatro vértices, expresadas en el sistema ", {}),
      ("Gauss-Krüger POSGAR 2007, Faja 2, meridiano central 69° Oeste (EPSG 5344)", {"b": True}),
      (", se consignan en el Anexo I y se reproducen en el plano de ubicación que se acompaña.", {})])

# ======================= IV. SUPERFICIE =======================
heading("IV. SUPERFICIE Y UNIDADES DE MEDIDA")
para([("La superficie solicitada asciende a ", {}),
      ("SEIS MIL TREINTA Y SIETE HECTÁREAS CON CINCUENTA ÁREAS (6.037,50 ha)", {"b": True}),
      (". Siendo la unidad de medida de los permisos de exploración de quinientas (500) hectáreas conforme "
       "al artículo 29 del Código de Minería, y computándose toda fracción como unidad completa, el "
       "pedimento comprende ", {}),
      ("TRECE (13) UNIDADES DE MEDIDA", {"b": True}),
      (", cantidad que no excede el máximo de veinte (20) unidades por permiso previsto en la norma citada.", {})])

# ======================= V. PLAZO =======================
heading("V. PLAZO SOLICITADO")
para([("Conforme al artículo 30, primer párrafo, del Código de Minería, correspondiendo ciento cincuenta "
       "(150) días por la primera unidad de medida más cincuenta (50) días por cada unidad adicional, se "
       "solicita el permiso por el plazo de ", {}),
      ("SETECIENTOS CINCUENTA (750) DÍAS CORRIDOS", {"b": True}),
      (", que comenzará a correr treinta (30) días después del otorgamiento.", {})])
para([("Se deja constancia de que el área solicitada no se encuentra comprendida en ninguna de las zonas "
       "de temporada registradas por el Catastro Minero, por lo que el ", {}),
      ("período de trabajo anual admitido es de todo el año", {"b": True}),
      (", sin restricción estacional.", {})])

# ======================= VI. MINERALES =======================
heading("VI. CATEGORÍA DE LOS MINERALES A EXPLORAR")
para([("La exploración tendrá por objeto sustancias de la ", {}),
      ("PRIMERA CATEGORÍA", {"b": True}),
      (", en particular cobre, oro y plata, comprendidas en el artículo 3, inciso a), del Código de Minería. ", {}),
      ("[Si también se explorarán sustancias de segunda categoría, consignarlo expresamente aquí y marcar "
       "la casilla respectiva en el formulario oficial.]", {"c": ACC, "i": True})])

# ======================= VII. TERRENOS =======================
heading("VII. ESTADO DE LOS TERRENOS Y PROPIETARIO DEL SUELO")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.line_spacing = 1.25; p.paragraph_format.space_after = Pt(7)
run(p, "El terreno superficial comprendido en el área solicitada se encuentra en estado de ")
fill(p, "[erial / cultivado / cercado / edificado — según informe de dominio]")
run(p, ", no encontrándose afectado a sitio público, histórico ni religioso, ni a reserva natural. "
       "Su propietario es ")
fill(p, "[apellido y nombre / razón social]")
run(p, ", con domicilio en ")
fill(p, "[domicilio]")
run(p, ".")

para([("Para el supuesto de desconocerse el nombre o domicilio del propietario, desde ya solicito se me "
       "entreguen los oficios pertinentes a los fines de su individualización, comprometiéndome a "
       "diligenciarlos dentro del plazo de quince (15) días que fija el artículo 37 de la Ley N.º 688-M.", {})])

# ======================= VIII. PROGRAMA =======================
heading("VIII. PROGRAMA MÍNIMO DE TRABAJOS")
para([("Se acompaña, en el formulario oficial respectivo y como ", {}), ("Anexo II", {"b": True}),
      (" del presente, el programa mínimo de trabajos a realizar con la estimación de las inversiones "
       "proyectadas y la indicación de los elementos y equipos a utilizar, todo ello conforme al artículo "
       "25, cuarto párrafo, del Código de Minería.", {})])
para([("Asimismo asumo el compromiso de dejar instalados los trabajos de exploración descriptos en dicho "
       "programa dentro de los treinta (30) días de notificado el otorgamiento del permiso, y de comunicar "
       "en el mismo plazo la situación del emplazamiento en el terreno y la descripción de los trabajos, a "
       "los fines de la verificación por parte de esa Autoridad Minera (art. 45 Ley N.º 688-M y art. 30, "
       "tercer párrafo, del Código de Minería).", {})])

# ======================= IX. DDJJ =======================
heading("IX. DECLARACIÓN JURADA — ARTÍCULOS 29 Y 30 DEL CÓDIGO DE MINERÍA")
para([("Declaro bajo juramento no encontrarme comprendido en las prohibiciones establecidas en los "
       "artículos 29, segundo párrafo, y 30, quinto párrafo, del Código de Minería. En particular declaro "
       "que, computando los permisos otorgados a mi nombre, a mis socios y por interpósita persona en el "
       "territorio de esta Provincia, registro:", {})], after=5)

table([12.0, 5.0], [
    ["Concepto", {"t": "Cantidad declarada", "a": "c"}],
    ["Permisos de exploración otorgados (máximo legal: 20)", {"t": "[COMPLETAR]", "c": ACC, "a": "c"}],
    ["Unidades de medida otorgadas (máximo legal: 400)", {"t": "[COMPLETAR]", "c": ACC, "a": "c"}],
    ["Unidades que consume el presente pedimento", {"t": "13", "b": True, "a": "c"}],
])

para([("Declaro asimismo que no he sido titular, por mí, por mis socios ni por interpósita persona, de "
       "permiso de exploración alguno caducado sobre la misma zona o parte de ella dentro del año anterior "
       "a la fecha de esta presentación.", {})])
para([("Tomo conocimiento de que la falsedad de la presente declaración se penará con multa igual a la del "
       "artículo 26 del Código de Minería y con la consiguiente pérdida de todos los derechos que se "
       "hubiesen peticionado u obtenido, los que serán inscriptos como vacantes (art. 25, quinto párrafo, "
       "del Código de Minería).", {})])

# ======================= X. CANON =======================
heading("X. CANON Y TASAS")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.line_spacing = 1.25; p.paragraph_format.space_after = Pt(7)
run(p, "Se acompaña constancia del pago provisional del canon de exploración correspondiente a las trece "
       "(13) unidades de medida solicitadas, conforme al artículo 215, inciso 3, del Código de Minería, "
       "por la suma de ")
fill(p, "[$ IMPORTE SEGÚN RESOLUCIÓN VIGENTE]")
run(p, ", así como de la tasa de solicitud de exploración por ")
fill(p, "[$ importe]")
run(p, ". Se deja constancia de que el valor del canon se actualiza por resolución de la Secretaría de "
       "Minería de la Nación conforme la variación del Índice de Precios al Consumidor (art. 213 del "
       "Código de Minería), por lo que se abona el importe vigente a la fecha de presentación.")

# ======================= XI. DOCUMENTAL =======================
heading("XI. DOCUMENTACIÓN QUE SE ACOMPAÑA")
bullets([
    "Formulario oficial de Solicitud de Permiso de Exploración, por triplicado.",
    "Formulario oficial de Programa Mínimo de Exploración, por triplicado (Anexo II).",
    "Plano de ubicación del área con coordenadas de los vértices (Anexo I).",
    "Constancia de pago del canon de exploración.",
    "Constancia de pago de la tasa de solicitud de exploración.",
    "Copia del documento de identidad, testimonio de poder o contrato social, según corresponda.",
])

# ======================= XII. PETITORIO =======================
heading("XII. PETITORIO")
para("Por todo lo expuesto, de V.S. solicito:", after=5)
bullets([
    "Me tenga por presentado, por parte y por constituido el domicilio legal indicado.",
    "Tenga por acompañada la documentación detallada en el Capítulo XI y por abonado el canon de exploración.",
    "Ordene el pase de las actuaciones al Registro Catastral a los fines de la asignación de matrícula "
    "catastral y del informe sobre ubicación, superficie y eventuales superposiciones (art. 39 Ley N.º 688-M).",
    "Ordene la anotación del pedimento en el Registro de Exploraciones y la publicación de edictos por dos (2) "
    "veces alternadas durante diez (10) días en el Boletín Oficial, a mi costa (art. 41 Ley N.º 688-M).",
    "Oportunamente, previo cumplimiento de los trámites de ley y no mediando oposición, otorgue el permiso "
    "de exploración solicitado por el plazo de setecientos cincuenta (750) días.",
])
para([("Proveer de conformidad,", {}), ("  SERÁ JUSTICIA.", {"b": True})], align="left", before=8, after=42)

# firmas
t = doc.add_table(rows=1, cols=2)
t.autofit = False
for i, txt in enumerate(("Firma del solicitante", "Firma y sello del profesional patrocinante")):
    c = t.rows[0].cells[i]
    c.width = Cm(8.5)
    tcPr = c._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single"); top.set(qn("w:sz"), "6"); top.set(qn("w:color"), "808080")
    borders.append(top)
    tcPr.append(borders)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    run(p, txt, size=9, color=GRY)

# ======================= ANEXO I =======================
para([("ANEXO I — COORDENADAS DEL ÁREA SOLICITADA", {"b": True})], align="center", after=3, new_page=True)
para([("Sistema Gauss-Krüger POSGAR 2007 · Faja 2 · meridiano central 69° Oeste · EPSG 5344",
       {"s": 9, "c": GRY})], align="center", after=12)

table([1.5, 1.9, 3.9, 3.9, 2.9, 2.9], [
    ["Vért.", "Esquina", {"t": "X (Este)", "a": "r"}, {"t": "Y (Norte)", "a": "r"},
     {"t": "Latitud", "a": "r"}, {"t": "Longitud", "a": "r"}],
    [{"t": "V1", "b": True}, "N.E.", {"t": "2.535.747,39", "a": "r"}, {"t": "6.676.339,31", "a": "r"},
     {"t": "−30,04922", "a": "r"}, {"t": "−68,62933", "a": "r"}],
    [{"t": "V2", "b": True}, "S.E.", {"t": "2.535.747,39", "a": "r"}, {"t": "6.664.839,31", "a": "r"},
     {"t": "−30,15295", "a": "r"}, {"t": "−68,62894", "a": "r"}],
    [{"t": "V3", "b": True}, "S.O.", {"t": "2.530.497,39", "a": "r"}, {"t": "6.664.839,31", "a": "r"},
     {"t": "−30,15310", "a": "r"}, {"t": "−68,68344", "a": "r"}],
    [{"t": "V4", "b": True}, "N.O.", {"t": "2.530.497,39", "a": "r"}, {"t": "6.676.339,31", "a": "r"},
     {"t": "−30,04936", "a": "r"}, {"t": "−68,68377", "a": "r"}],
])

table([10.0, 7.0], [
    ["Dato", "Valor"],
    ["Lado Este–Oeste", "5,25 km"],
    ["Lado Norte–Sur", "11,50 km"],
    ["Superficie total", {"t": "6.037,50 hectáreas", "b": True}],
    ["Unidades de medida (500 ha; fracción = unidad completa)", {"t": "13", "b": True}],
    ["Plazo que corresponde (150 + 50 × 12)", {"t": "750 días corridos", "b": True}],
    ["Departamento", "Jáchal"],
    ["Matrícula catastral", {"t": "la asigna el Registro Catastral", "c": ACC}],
])

para([("Verificación previa realizada sobre el padrón del Catastro Minero Digital (servicio WFS oficial): "
       "el área no presenta superposición con permisos de exploración, manifestaciones, minas, solicitudes, "
       "canteras ni áreas de aprovechamiento común. La distancia mínima al derecho vigente más próximo es "
       "de doscientos cincuenta y dos (252) metros. Sin perjuicio de ello, la superficie libre disponible "
       "será la que determine el Registro Catastral en su informe.", {"s": 9})])

# ======================= ANEXO II =======================
para([("ANEXO II — PROGRAMA MÍNIMO DE TRABAJOS E INVERSIONES", {"b": True})], align="center", after=3, new_page=True)
para([("Artículo 25, cuarto párrafo, del Código de Minería", {"s": 9, "c": GRY})], align="center", after=12)

para([("1. Trabajos de prospección", {"b": True})], align="left", after=5)
table([9.5, 3.0, 4.5], [
    ["Tarea", {"t": "¿Se ejecuta?", "a": "c"}, {"t": "Plazo estimado", "a": "c"}],
    ["Interpretación de imágenes satelitales", {"t": "Sí", "a": "c"}, {"t": "[completar]", "c": ACC, "a": "c"}],
    ["Relevamiento geológico de superficie", {"t": "Sí", "a": "c"}, {"t": "[completar]", "c": ACC, "a": "c"}],
    ["Muestreo geoquímico", {"t": "Sí", "a": "c"}, {"t": "[completar]", "c": ACC, "a": "c"}],
    ["Calicatas y extracción de muestras", {"t": "Sí", "a": "c"}, {"t": "[completar]", "c": ACC, "a": "c"}],
    ["Relevamiento topográfico", {"t": "[definir]", "c": ACC, "a": "c"}, {"t": "[completar]", "c": ACC, "a": "c"}],
    ["Prospección geofísica", {"t": "[definir]", "c": ACC, "a": "c"}, {"t": "[completar]", "c": ACC, "a": "c"}],
], size=9, gap=1)

para([("2. Trabajos de exploración", {"b": True})], align="left", after=5)
table([9.5, 3.0, 4.5], [
    ["Tarea", {"t": "¿Se ejecuta?", "a": "c"}, {"t": "Plazo estimado", "a": "c"}],
    ["Laboreos mineros", {"t": "[definir]", "c": ACC, "a": "c"}, {"t": "[completar]", "c": ACC, "a": "c"}],
    ["Sondajes", {"t": "[definir]", "c": ACC, "a": "c"}, {"t": "[completar]", "c": ACC, "a": "c"}],
    ["Otros métodos", {"t": "[definir]", "c": ACC, "a": "c"}, {"t": "[completar]", "c": ACC, "a": "c"}],
], size=9, gap=1)

para([("3. Maquinaria, equipo y personal", {"b": True})], align="left", after=5)
table([6.5, 4.5, 3.0, 3.0], [
    ["Elemento", "Capacidad / clase", {"t": "Cantidad", "a": "c"}, {"t": "¿Propio?", "a": "c"}],
    ["Compresores", {"t": "[completar]", "c": ACC}, {"t": "—", "c": ACC, "a": "c"}, {"t": "—", "c": ACC, "a": "c"}],
    ["Perforadoras", {"t": "[completar]", "c": ACC}, {"t": "—", "c": ACC, "a": "c"}, {"t": "—", "c": ACC, "a": "c"}],
    ["Vehículos", {"t": "[completar]", "c": ACC}, {"t": "—", "c": ACC, "a": "c"}, {"t": "—", "c": ACC, "a": "c"}],
    ["Operarios", "—", {"t": "[completar]", "c": ACC, "a": "c"}, {"t": "—", "a": "c"}],
    ["Técnicos y profesionales", "—", {"t": "[completar]", "c": ACC, "a": "c"}, {"t": "—", "a": "c"}],
], size=9, gap=1)

para([("4. Estimación de inversiones", {"b": True})], align="left", after=5)
table([11.5, 5.5], [
    ["Rubro", {"t": "Monto estimado", "a": "r"}],
    ["Construcciones e instalaciones", {"t": "[completar]", "c": ACC, "a": "r"}],
    ["Caminos y huellas", {"t": "[completar]", "c": ACC, "a": "r"}],
    ["Laboreos y perforaciones", {"t": "[completar]", "c": ACC, "a": "r"}],
    ["Maquinaria y equipos", {"t": "[completar]", "c": ACC, "a": "r"}],
    ["Gastos generales", {"t": "[completar]", "c": ACC, "a": "r"}],
    [{"t": "TOTAL", "b": True}, {"t": "[completar]", "c": ACC, "a": "r", "b": True}],
], size=9, gap=1)

para([("5. Protección ambiental", {"b": True})], align="left", after=5)
para([("El Informe de Impacto Ambiental se presentará por separado, con arreglo a lo dispuesto por el "
       "artículo 251 y concordantes del Código de Minería (Título XIII, Sección Segunda) y el artículo 34 "
       "de la Ley N.º 688-M, con carácter previo al inicio de toda actividad, incluida la prospección. "
       "Rige asimismo la Ley Provincial N.º 2076-L en cuanto a las sustancias prohibidas en los procesos "
       "de lixiviación de minerales metalíferos.", {"s": 10})])

para([("6. Compromiso de información técnica", {"b": True})], align="left", after=5)
para([("De conformidad con lo dispuesto por el artículo 30, última parte, del Código de Minería y el "
       "artículo 45 de la Ley N.º 688-M, asumo el compromiso de presentar copia de la información y de la "
       "documentación técnica obtenida en el curso de las investigaciones dentro de los noventa (90) días "
       "de vencido el permiso, bajo pena de una multa igual al doble del canon abonado.", {"s": 10})])

# ======================= ANEXO III =======================
para([("ANEXO III — CRONOGRAMA PROCESAL Y PLAZOS FATALES", {"b": True})], align="center", after=3, new_page=True)
para([("Hoja de control interna — no forma parte del escrito a presentar", {"s": 9, "c": ACC})],
     align="center", after=10)
para([("Fechas calculadas para una presentación el 7 de septiembre de 2026. Los plazos de la Ley N.º 688-M "
       "se computan en días hábiles y los del Código de Minería en días corridos (art. 26 Ley N.º 688-M). "
       "Las fechas de trámite interno son estimaciones operativas.", {"s": 9})], after=8)

table([2.6, 6.6, 3.9, 3.9], [
    ["Fecha", "Hito", "Norma", "Carácter"],
    ["28/09/2026", "Publicar edictos — 2 veces alternadas en 10 días", "art. 41 L. 688-M", "Carga del solicitante"],
    ["26/10/2026", "Acreditar la publicación (20 días hábiles)", "art. 41 L. 688-M", "Carga del solicitante"],
    ["28/10/2026", "Cierre del plazo de oposiciones", "art. 42 L. 688-M", "Plazo de terceros"],
    ["25/01/2027", "Otorgamiento estimado del permiso", "art. 29 inc. c L. 688-M", "Estimación"],
    [{"t": "24/02/2027", "c": RED, "b": True}, {"t": "Instalar los trabajos del programa mínimo", "b": True},
     "art. 45 L. 688-M; art. 41 inc. a CM", {"t": "FATAL — revocación", "c": RED, "b": True}],
    [{"t": "21/12/2027", "c": RED, "b": True},
     {"t": "1.ª liberación: desafectar 2.018,8 ha — retiene 4.018,8 ha", "b": True},
     "art. 30 párr. 2 CM", {"t": "FATAL — libera Catastro + multa", "c": RED, "b": True}],
    [{"t": "24/01/2029", "c": RED, "b": True},
     {"t": "2.ª liberación: desafectar 1.009,4 ha — retiene 3.009,4 ha", "b": True},
     "art. 30 párr. 2 CM", {"t": "FATAL — libera Catastro + multa", "c": RED, "b": True}],
    [{"t": "15/03/2029", "c": RED, "b": True},
     {"t": "Vencimiento del permiso (750 días corridos)", "b": True},
     "art. 30 CM; art. 48 L. 688-M", {"t": "FATAL — caducidad de pleno derecho", "c": RED, "b": True}],
    ["13/06/2029", "Presentar la información técnica", "art. 30 in fine CM", "Multa: doble del canon"],
])

para([("Advertencia sobre las liberaciones. ", {"b": True}),
      ("El permiso pierde la mitad de su superficie por ministerio de la ley. La decisión relevante no es "
       "si liberar, sino qué mitad conservar, y debe fundarse en los resultados de la prospección de los "
       "primeros diez meses. La petición debe presentarse ", {}),
      ("antes", {"i": True}),
      (" del cumplimiento del plazo, indicando las coordenadas de cada vértice del área que se mantiene; "
       "en su defecto, el Catastro Minero libera las zonas a su criterio y se aplica una multa igual al "
       "canon abonado.", {})], after=6)

# ======================= ANEXO IV =======================
para([("ANEXO IV — CONTROL PREVIO A LA PRESENTACIÓN", {"b": True})], align="center", after=3, new_page=True)
para([("Hoja de control interna — no forma parte del escrito a presentar", {"s": 9, "c": ACC})],
     align="center", after=12)

table([1.2, 11.3, 4.5], [
    ["", "Recaudo", "Fundamento"],
    [{"t": "☐", "a": "c"}, "Informe de área libre solicitado a Catastro Minero sobre las coordenadas del Anexo I", "art. 13 L. 688-M"],
    [{"t": "☐", "a": "c"}, "Informe de dominio del terreno superficial (propietario y estado de los terrenos)", "art. 25 CM"],
    [{"t": "☐", "a": "c"}, "Datos personales completos y domicilio legal constituido en el radio de la Autoridad", "art. 8 L. 688-M"],
    [{"t": "☐", "a": "c"}, "Control de topes verificado, incluyendo socios e interpósita persona", "art. 29 CM"],
    [{"t": "☐", "a": "c"}, "Canon vigente confirmado en Escribanía de Minas", "arts. 213 y 215 CM"],
    [{"t": "☐", "a": "c"}, "Constancia de pago del canon obtenida (sin ella: rechazo sin recurso)", "art. 25 CM"],
    [{"t": "☐", "a": "c"}, "Programa mínimo de trabajos completo, en formulario oficial", "art. 25 CM"],
    [{"t": "☐", "a": "c"}, "Plano de ubicación con coordenadas de los vértices", "art. 19 CM"],
    [{"t": "☐", "a": "c"}, "Todo impreso por triplicado en formulario oficial, sin abreviaturas ni raspaduras", "arts. 8 y 10 L. 688-M"],
    [{"t": "☐", "a": "c"}, "Presentación en Mesa de Entradas con cargo del Escribano de Minas (fija la prioridad)", "arts. 11 y 35 L. 688-M"],
])

para([("Si el Registro Catastral informa superposición parcial", {"b": True}),
      (", se corre vista por cinco (5) días hábiles; el silencio se tiene por aceptación y el trámite "
       "prosigue por la parte libre remanente (art. 40 Ley N.º 688-M). Si la superposición fuera total, la "
       "solicitud se desestima y el canon se reintegra dentro de los diez (10) días.", {})], after=6)

doc.core_properties.title = "Solicitud de Permiso de Exploración — Jáchal"
doc.core_properties.author = "Boletín Minero San Juan"
doc.save("Solicitud-Permiso-Exploracion-Jachal.docx")
print("OK — Solicitud-Permiso-Exploracion-Jachal.docx")
